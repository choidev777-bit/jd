"""자소서 최종본(final.md 또는 draft-vN.md)을 Word 파일로 변환한다.

사용법:
    python md_to_docx.py <final.md> [<출력.docx>]

출력 경로를 생략하면 같은 폴더에 같은 이름의 .docx로 저장한다.

변환 규칙:
- "# 제목" 줄은 문서 제목.
- "## N. 제목 (제한)" 문항 헤더는 굵게. 괄호의 글자 수 제한은 뺀다.
- "> 문항: ..." 줄은 기울임으로 문항 원문을 넣는다.
- 그 외 ">" 줄(보고줄)과 문항 앞의 "- 작성일:" 같은 메타 줄은 뺀다.
- 본문은 빈 줄 기준으로 단락을 나눈다.

필요 패키지: python-docx  (없으면: python -m pip install python-docx)
"""
import os
import sys

from count_chars import HEADER_RE

QUESTION_PREFIX = "> 문항:"


def parse_for_docx(text):
    title = None
    sections = []
    current = None
    buf = []

    def flush():
        if current is not None and buf:
            current["paragraphs"].append(" ".join(buf).strip())
            buf.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if current is None and stripped.startswith("# ") and title is None:
            title = stripped[2:].strip()
            continue
        m = HEADER_RE.match(line)
        if m:
            flush()
            current = {"heading": f"{m.group(1)}. {m.group(2).strip()}",
                       "question": None, "paragraphs": []}
            sections.append(current)
            continue
        if current is None:
            continue
        if stripped.startswith(QUESTION_PREFIX):
            current["question"] = stripped[len(QUESTION_PREFIX):].strip()
            continue
        if stripped.startswith(">"):
            continue
        if not stripped:
            flush()
            continue
        buf.append(stripped)
    flush()
    return {"title": title or "자기소개서", "sections": sections}


def build_docx(doc, out_path):
    from docx import Document
    from docx.shared import Pt

    d = Document()
    style = d.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)

    p = d.add_paragraph()
    r = p.add_run(doc["title"])
    r.bold = True
    r.font.size = Pt(16)

    for s in doc["sections"]:
        d.add_paragraph()
        h = d.add_paragraph()
        hr = h.add_run(s["heading"])
        hr.bold = True
        hr.font.size = Pt(13)
        if s["question"]:
            q = d.add_paragraph()
            qr = q.add_run(s["question"])
            qr.italic = True
        for para in s["paragraphs"]:
            d.add_paragraph(para)

    d.save(out_path)
    return out_path


def main(argv):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(argv) not in (2, 3):
        print(__doc__)
        return 2
    src = argv[1]
    out = argv[2] if len(argv) == 3 else os.path.splitext(src)[0] + ".docx"
    try:
        import docx  # noqa: F401
    except ImportError:
        print("python-docx가 없습니다. 실행: python -m pip install python-docx")
        return 1
    with open(src, encoding="utf-8") as f:
        doc = parse_for_docx(f.read())
    if not doc["sections"]:
        print("문항 헤더(## N. 제목)를 찾지 못했습니다.")
        return 1
    build_docx(doc, out)
    print(f"저장: {out} (문항 {len(doc['sections'])}개)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

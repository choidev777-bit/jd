import os
import tempfile
import unittest

from docx import Document

from md_to_docx import parse_for_docx, build_docx


SAMPLE = """# 테스트회사 개발 자소서 v2

- 작성일: 2026-09-05
- 공고: posting.md
- 이전 버전 대비 변경: 1번 축약

## 1. 지원동기 (1,000자 이내)
> 문항: 테스트회사에 지원한 이유를 적어 주세요.
첫 문단입니다.

둘째 문단입니다.
> 공백 포함 20자 / 공백 제외 18자 · 사용 경험: 인턴

## 2. 자유기술
본문만 있는 문항.
"""


class ParseTest(unittest.TestCase):
    def test_extracts_title_sections_question_and_paragraphs(self):
        doc = parse_for_docx(SAMPLE)
        self.assertEqual(doc["title"], "테스트회사 개발 자소서 v2")
        self.assertEqual(len(doc["sections"]), 2)
        s1 = doc["sections"][0]
        self.assertEqual(s1["heading"], "1. 지원동기")
        self.assertEqual(s1["question"], "테스트회사에 지원한 이유를 적어 주세요.")
        self.assertEqual(s1["paragraphs"], ["첫 문단입니다.", "둘째 문단입니다."])
        s2 = doc["sections"][1]
        self.assertEqual(s2["heading"], "2. 자유기술")
        self.assertIsNone(s2["question"])
        self.assertEqual(s2["paragraphs"], ["본문만 있는 문항."])

    def test_meta_and_report_lines_are_dropped(self):
        doc = parse_for_docx(SAMPLE)
        all_text = " ".join(p for s in doc["sections"] for p in s["paragraphs"])
        self.assertNotIn("작성일", all_text)
        self.assertNotIn("공백 포함", all_text)


class BuildTest(unittest.TestCase):
    def test_docx_has_bold_heading_italic_question_plain_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "final.docx")
            build_docx(parse_for_docx(SAMPLE), out)
            d = Document(out)
            texts = [p.text for p in d.paragraphs if p.text.strip()]
            self.assertEqual(texts[0], "테스트회사 개발 자소서 v2")
            self.assertIn("1. 지원동기", texts)
            self.assertIn("테스트회사에 지원한 이유를 적어 주세요.", texts)
            self.assertIn("첫 문단입니다.", texts)
            self.assertNotIn("공백 포함 20자 / 공백 제외 18자 · 사용 경험: 인턴", texts)
            heading = next(p for p in d.paragraphs if p.text == "1. 지원동기")
            self.assertTrue(all(r.bold for r in heading.runs))
            question = next(p for p in d.paragraphs if p.text.startswith("테스트회사에 지원한"))
            self.assertTrue(all(r.italic for r in question.runs))
            body = next(p for p in d.paragraphs if p.text == "첫 문단입니다.")
            self.assertFalse(any(r.bold or r.italic for r in body.runs))


if __name__ == "__main__":
    unittest.main()
